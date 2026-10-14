# -*- coding: utf-8 -*-
"""Genera manual_tecnico.docx y manual_usuario.docx en la raiz del proyecto.

Los documentos describen el funcionamiento ACTUAL del sistema RCS.
Todo el contenido (constantes, firmas, rutas, mensajes) fue verificado contra
el codigo fuente antes de redactarse.
"""
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
VERSION = "Versión 1.0 — 28 de septiembre de 2026"


# --------------------------------------------------------------------------
# Utilidades de formato
# --------------------------------------------------------------------------
def base_document():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    estilo.paragraph_format.space_after = Pt(6)
    return doc


def titulo(doc, texto, sub=None):
    h = doc.add_heading(texto, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    if sub:
        p = doc.add_paragraph(sub)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.italic = True
            r.font.size = Pt(10)
    p = doc.add_paragraph(VERSION)
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    doc.add_paragraph()


def seccion(doc, texto):
    doc.add_heading(texto, level=1)


def subseccion(doc, texto):
    doc.add_heading(texto, level=2)


def parrafo(doc, texto):
    doc.add_paragraph(texto)


def vineta(doc, texto):
    doc.add_paragraph(texto, style="List Bullet")


def codigo(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(texto)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    return p


def tabla(doc, cabeceras, filas):
    n = len(cabeceras)
    for i, fila in enumerate(filas):
        if len(fila) != n:
            raise ValueError(
                "Fila %d de la tabla '%s': %d valores para %d columnas."
                % (i, cabeceras[0], len(fila), n)
            )
    t = doc.add_table(rows=1, cols=n)
    try:
        t.style = "Light Grid Accent 1"
    except KeyError:
        t.style = "Table Grid"
    celdas = t.rows[0].cells
    for i, h in enumerate(cabeceras):
        celdas[i].text = ""
        r = celdas[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
    for fila in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(fila):
            celdas[i].text = ""
            r = celdas[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9)
    doc.add_paragraph()
    return t


# --------------------------------------------------------------------------
# MANUAL TÉCNICO
# --------------------------------------------------------------------------
def generar_tecnico():
    doc = base_document()
    titulo(
        doc,
        "Manual técnico — RCS",
        "Revisión de publicaciones sísmicas, detección de eventos no publicados "
        "y asignación de perfiles de subducción",
    )

    # ---------------------------------------------------------------- 1
    seccion(doc, "1. Propósito y alcance")
    parrafo(
        doc,
        "RCS es una aplicación de escritorio que procesa la salida de un "
        "sistema de detección sísmica (Seisan) y la contrasta con el catálogo público de "
        "eventos, con el fin de responder tres preguntas operativas: qué eventos del "
        "catálogo público no fueron publicados por el sistema, qué eventos se publicaron "
        "con datos distintos a los procesados, y qué eventos de la revisión manual tienen "
        "una localización dudosa respecto del perfil de subducción bajo ellos.",
    )
    parrafo(
        doc,
        "Este documento describe el funcionamiento del sistema tal como está implementado: "
        "su arquitectura, el flujo de cada etapa, los archivos que produce, los algoritmos "
        "que decide sus resultados y los parámetros que los gobiernan. Está dirigido a "
        "quien necesite mantener, ampliar o depurar el código.",
    )
    parrafo(doc, "El sistema se compone de:")
    vineta(doc, "Una interfaz gráfica (app.py) que orquesta la ejecución y muestra el registro y los mapas.")
    vineta(doc, "Un conjunto de módulos de proceso que se ejecutan como subprocesos independientes.")
    vineta(doc, "Dos scripts de shell para el arranque con entorno virtual y para la extracción remota de datos.")

    # ---------------------------------------------------------------- 2
    seccion(doc, "2. Arquitectura general")
    parrafo(
        doc,
        "app.py es una aplicación de escritorio construida con ttkbootstrap sobre Tkinter. "
        "No contiene lógica de negocio: su función es lanzar cada módulo como un subproceso, "
        "capturar su salida estándar para mostrarla en la pestaña de registro, interpretar las "
        "líneas de progreso que los módulos emiten, y recibir los resultados para incrustar "
        "los mapas en un panel. Todos los cálculos se ejecutan en los subprocesos.",
    )
    parrafo(
        doc,
        "Los subprocesos se lanzan con subprocess.Popen usando cwd=os.getcwd(), es decir, el "
        "directorio desde donde se invocó la aplicación. Esto tiene una consecuencia importante: "
        "las rutas relativas de datos e informes (rutas.py) se resuelven contra ese directorio, "
        "no contra la carpeta donde está el código. En cambio, los archivos de datos grandes "
        "—grillas, relieves, catálogo histórico, localidades— se localizan con rutas ancladas "
        "al propio archivo del módulo, de modo que siempre se encuentran.",
    )
    parrafo(
        doc,
        "Los módulos que importan prog (proc_query_harz_2, compara, generajson, repetidos, "
        "repetidosexclu) emiten una línea de progreso por etapa con el formato "
        "@@PROGRESO@@ <fracción>. app.py la reconoce, la convierte a progreso global "
        "multiplicando por el peso de la etapa activa y actualiza una barra monótona de 0 a "
        "1000. Los módulos que no importan prog simplemente no aportan progreso, y la barra "
        "avanza por tramos al cambiar de etapa.",
    )
    parrafo(
        doc,
        "La emisión de progreso se activa únicamente cuando existe la variable de entorno "
        "RV_PROG, que app.py define en cada subproceso. Al ejecutar un módulo directamente "
        "desde la línea de comandos, el progreso no se imprime y la salida es limpia.",
    )

    subseccion(doc, "2.1 Inventario de módulos")
    tabla(
        doc,
        ["Módulo", "Responsabilidad", "Emite progreso"],
        [
            ("app.py", "Interfaz gráfica: orquesta las etapas, muestra el registro y los mapas.", "—"),
            ("rutas.py", "Rutas de datos e informes; crea las carpetas de salida.", "—"),
            ("prog.py", "Protocolo de emisión de progreso por subproceso.", "—"),
            ("proc_query_harz_2.py", "Convierte el CSV de la consulta pública en CSV con formato Seisan y un .dat legible.", "Sí"),
            ("revisaselect.py", "Parsea select.out y separa eventos válidos de excluidos.", "No"),
            ("revisacollect.py", "Construye el CSV maestro de la revisión y los CSV por analista.", "No"),
            ("compara.py", "Compara el catálogo público contra el select: no publicados y no actualizados.", "Sí"),
            ("revisaexcluidos.py", "Normaliza el listado de excluidos y detecta líneas anómalas.", "No"),
            ("repetidosexclu.py", "Busca repeticiones entre los excluidos y el select completo.", "Sí"),
            ("repetidos.py", "Detecta repeticiones dentro de cada fuente por separado.", "Sí"),
            ("generajson.py", "Asigna perfiles, clasifica sospechosos y escribe el JSON y los CSV de salida.", "Sí"),
            ("asigna_perfiles.py", "Detecta los perfiles de grillas y asigna un perfil a cada evento.", "No"),
            ("sismicidad.py", "Clasifica un evento como posible mal localizado.", "No"),
            ("plotear.py", "Dibuja los mapas y el panel de análisis; gestiona la interactividad.", "No"),
        ],
    )
    tabla(
        doc,
        ["Script de apoyo", "Función"],
        [
            ("supervisor.sh", "Crea el entorno virtual, instala las dependencias y lanza app.py."),
            ("extraesc.sh", "Extrae eventos del servidor SeisComp por scquery y los convierte a CSV."),
            ("traequery.py", "Descarga del catálogo público por eventquery, de forma interactiva."),
        ],
    )

    # ---------------------------------------------------------------- 3
    seccion(doc, "3. Estructura de directorios y archivos")
    subseccion(doc, "3.1 Carpetas de trabajo (rutas.py)")
    parrafo(
        doc,
        "rutas.py define cinco carpetas de salida, relativas al directorio de trabajo. Sus "
        "funciones p_datos, p_informes, p_analistas, p_ploteo y p_trabajo crean la carpeta de "
        "forma perezosa la primera vez que se usan, de modo que no hace falta preparar la "
        "estructura antes de empezar. La función asegurar_dirs existe pero no la invoca la "
        "aplicación.",
    )
    tabla(
        doc,
        ["Carpeta", "Contenido"],
        [
            ("datos/", "CSV y JSON de trabajo: select normalizado, eventos con perfil, conteos, sospechosos."),
            ("informes/", "Reportes de texto de comparación, repetidos y exclusiones."),
            ("analistas/", "Un CSV por analista con los eventos que le corresponden."),
            ("ploteo/", "Bitácora de la sesión: percibidos.txt y los registros de resumen."),
            ("trabajo/", "Archivos intermedios, que se renombran o se borran al terminar cada etapa."),
        ],
    )

    subseccion(doc, "3.2 Archivos de entrada y de datos grandes")
    parrafo(
        doc,
        "Estos archivos se localizan con rutas ancladas al propio módulo, por lo que no "
        "dependen del directorio de trabajo:",
    )
    tabla(
        doc,
        ["Archivo", "Quién lo usa", "Función"],
        [
            ("grillas/", "asigna_perfiles.py, plotear.py", "Perfiles de subducción: pares slabP###.tmp y topoP###.tmp."),
            ("select.out", "revisaselect.py", "Salida del sistema de detección; define el universo de eventos revisados."),
            ("base_2020_2026.dat", "sismicidad.py, plotear.py", "Catálogo histórico de sismicidad, solo como referencia."),
            ("localidades.csv", "plotear.py", "Puntos de referencia para rotular los mapas."),
            ("relieve_chile.tif", "plotear.py", "Relieve regional; tiene prioridad sobre el relieve global."),
            ("NE2_LR_LC_SR_W_DR.tif", "plotear.py", "Relieve global de respaldo, para áreas fuera del recorte regional."),
        ],
    )
    parrafo(
        doc,
        "El CSV de entrada de la aplicación se indica en el primer argumento de app.py y su "
        "contenido determina el nombre base de varios archivos intermedios. El segundo "
        "argumento es un archivo .dat temporal que app.py elimina al terminar el análisis.",
    )

    # ---------------------------------------------------------------- 4
    seccion(doc, "4. Ejecución del sistema")
    subseccion(doc, "4.1 Arranque con entorno virtual")
    parrafo(
        doc,
        "supervisor.sh es el punto de entrada recomendado. Si el entorno virtual no existe, "
        "lo crea e instala las dependencias de requirements.txt; después ejecuta app.py "
        "pasándole los argumentos sin modificar.",
    )
    codigo(doc, "./supervisor.sh <archivo_entrada.csv> <archivo_salida.dat>")
    parrafo(
        doc,
        "Shapely se instala deliberadamente compilado desde fuente (--no-binary shapely). "
        "El wheel precompilado trae su propia copia de GEOS y provoca un fallo de segmentación "
        "al usarlo junto con cartopy, que enlaza la GEOS del sistema. Compilarlo requiere "
        "libgeos-dev y un compilador C en la máquina; si falta, supervisor.sh lo indica "
        "explícitamente.",
    )

    subseccion(doc, "4.2 Línea de comandos de cada módulo")
    parrafo(
        doc,
        "Todos los módulos se pueden ejecutar de forma independiente, lo que resulta útil "
        "para depurar una etapa concreta. Ninguno usa argparse: los argumentos son posicionales "
        "y las opciones se escriben como --clave=valor.",
    )
    tabla(
        doc,
        ["Comando", "Efecto"],
        [
            ("python3 app.py <entrada.csv> <salida.dat>", "Abre la interfaz gráfica y guarda el directorio de trabajo actual."),
            ("python3 generajson.py <csv> <fuente> [--umbral=KM] [--k=K] [--umbral-perp=KM] [--margen-borde=KM]", "Genera el JSON, el CSV de sospechosos y el conteo de perfiles. La fuente es seisan, eventquery o nopub."),
            ("python3 plotear.py <eventos_<fuente>.json> <fuente>", "Abre los mapas de una fuente ya generada."),
            ("python3 asigna_perfiles.py --listar", "Lista los perfiles detectados en grillas/."),
            ("python3 asigna_perfiles.py <lon> <lat> <prof>", "Asigna un perfil a una coordenada concreta e imprime el resultado."),
            ("python3 sismicidad.py <eventos_<fuente>.json>", "Indica, evento por evento, si es sospechoso."),
            ("python3 compara.py <csv_publico> <csv_seisan>", "Compara ambas fuentes y escribe los reportes."),
            ("python3 repetidos.py <csv_publico> <csv_seisan>", "Detecta repeticiones dentro de cada fuente."),
            ("python3 repetidosexclu.py <csv_excluidos> <csv_select>", "Busca repeticiones entre excluidos y select."),
            ("python3 traequery.py", "Descarga eventos del catálogo público por SSH, de forma interactiva."),
        ],
    )
    parrafo(
        doc,
        "Advertencia sobre las rutas: como los módulos resuelven datos/, informes/ y trabajo/ "
        "contra el directorio actual, deben ejecutarse desde la carpeta del proyecto. Si se "
        "ejecutan desde otro sitio, los archivos se crearán allí y no se encontrarán los datos "
        "ya calculados.",
    )

    # ---------------------------------------------------------------- 5
    seccion(doc, "5. Pipeline de análisis")
    parrafo(
        doc,
        "El botón de análisis ejecuta una secuencia fija de etapas. La barra de progreso "
        "reparte el avance entre ellas según los pesos declarados en ANALISIS_PLAN, y la etapa "
        "activa se determina comparando el nombre del script que se está lanzando con el campo "
        "scripts de cada entrada del plan.",
    )
    tabla(
        doc,
        ["#", "Etiqueta en pantalla", "Script", "Peso", "Qué produce"],
        [
            ("1", "Procesando query", "proc_query_harz_2.py", "10 %", "datos/new_<base>.csv, datos/new_2_<base>.csv y el .dat"),
            ("2", "Revisando select", "revisaselect.py", "15 %", "trabajo/newcollect.txt, informes/excluidos.txt"),
            ("3", "Revisando collect", "revisacollect.py", "15 %", "datos/salida_collect.csv y los CSV por analista"),
            ("4", "Comparando publicados v/s procesados", "compara.py", "25 %", "Los cuatro reportes de no publicados y no actualizados"),
            ("5", "Revisando excluidos", "revisaexcluidos.py, repetidosexclu.py", "10 %", "datos/excluidos.csv e informes/rep_seisan_exclu.txt"),
            ("6", "Revisando repetidos", "repetidos.py", "25 %", "Los cuatro reportes de repetidos"),
        ],
    )
    parrafo(doc, "Detalles del flujo que conviene conocer:")
    vineta(
        doc,
        "El paso 4 solo se ejecuta si datos/salida_collect.csv existe y no está vacío; en "
        "caso contrario se omite y el registro lo indica.",
    )
    vineta(
        doc,
        "El paso 5 es condicional en su segunda parte: si no hay exclusiones, app.py borra los "
        "residuos de una corrida anterior y anota en el registro que no se excluyeron eventos.",
    )
    vineta(doc, "Al terminar, app.py elimina el archivo .dat recibido como segundo argumento.")
    parrafo(
        doc,
        "La barra de progreso es monótona: solo sube. Esto evita el efecto de retroceso "
        "visible cuando una etapa reporta una fracción menor que la ya alcanzada.",
    )

    # ---------------------------------------------------------------- 6
    seccion(doc, "6. Pipeline de fuentes")
    parrafo(
        doc,
        "Terminado el análisis, los botones de fuente generan el JSON de eventos con su perfil "
        "asignado y abren el panel de mapas. Los tres comandos difieren solo en el CSV de "
        "partida y en el nombre de la fuente, que a su vez determina el nombre de los archivos "
        "de salida.",
    )
    tabla(
        doc,
        ["Botón", "Comando", "Salida"],
        [
            ("Datos de Seisan", "generajson.py datos/salida_collect.csv seisan", "eventos_seisan.json"),
            ("Datos de eventquery", "concatena datos/new_2_*.csv y ejecuta generajson.py datos/todos_eventquery.csv eventquery", "eventos_eventquery.json"),
            ("No publicados de Seisan", "generajson.py datos/no_pub_desde_2_5_estricto.csv seisan", "Sobrescribe eventos_seisan.json"),
        ],
    )
    parrafo(
        doc,
        "Dos advertencias importantes. Primera: la fuente no publicados usa el nombre seisan, "
        "por lo que sobrescribe el JSON de la fuente de Seisan; conviene revisar los no "
        "publicados antes de generar sobre ellos. Segunda: la concatenación de eventquery "
        "suma varios archivos de entrada en uno solo, de modo que el número de eventos del "
        "catálogo público puede superar al de una corrida individual.",
    )

    # ---------------------------------------------------------------- 7
    seccion(doc, "7. Asignación de perfiles")
    parrafo(
        doc,
        "Este es el núcleo del sistema y la parte con más sutilezas. Su responsabilidad es "
        "decidir, para cada evento, a qué perfil de subducción pertenece y con qué "
        "confianza, o bien declarar que ningún perfil del conjunto lo cubre.",
    )

    subseccion(doc, "7.1 Formato de los archivos de grillas")
    parrafo(
        doc,
        "Cada perfil se compone de un par de archivos en la carpeta grillas/: uno con la "
        "geometría de la placa y la profundidad de la interfase, y otro con la topografía. Son "
        "tablas de texto separadas por tabulador, de siete columnas numéricas, sin cabecera, "
        "en formato GMT. Se pueden admitir líneas de comentario iniciadas por #.",
    )
    tabla(
        doc,
        ["Columna", "Símbolo", "Contenido", "Uso en el código"],
        [
            ("0", "p", "Distancia a lo largo del perfil, en km desde el inicio", "Eje del perfil; define la línea central y la interpolación de la profundidad."),
            ("1", "q", "Distancia perpendicular a la línea central, en km", "Informativo; la línea central se reconstruye agrupando por p."),
            ("2", "s", "Longitud del punto", "Geometría de la sección."),
            ("3", "r", "Latitud del punto", "Geometría de la sección y cálculo de la latitud media."),
            ("4", "x", "Longitud del nodo de grilla", "No se usa directamente."),
            ("5", "y", "Latitud del nodo de grilla", "No se usa directamente."),
            ("6", "z", "En slab: profundidad de la interfase, en km y con signo negativo. En topo: altitud, en metros", "Modelo de profundidad del slab; relieve del perfil."),
        ],
    )
    parrafo(
        doc,
        "Conviene aclarar que estos archivos no son grillas regulares sino conjuntos de puntos "
        "muestreados: el número de filas y la disposición de las columnas varían entre perfiles. "
        "La profundidad puede venir como NaN en los extremos del perfil, y el código la ignora.",
    )

    subseccion(doc, "7.2 Detección de perfiles")
    parrafo(
        doc,
        "detectar_perfiles() recorre grillas/ y reconoce exclusivamente los archivos que "
        "cumplen el patrón slabP<dígitos>.tmp, ordenándolos por número. Para cada uno, intenta "
        "cargar la topografía con el mismo ancho de dígitos que el slab y, si no la encuentra, "
        "recurre al formato de tres dígitos. Esta doble búsqueda existe porque los catálogos "
        "antiguos usan slabP01.tmp mientras que los actuales usan slabP001.tmp.",
    )
    parrafo(
        doc,
        "Como la identidad del perfil se deriva del número del nombre del archivo, dos "
        "conjuntos distintos pueden producir las mismas etiquetas. Si la carpeta contiene las "
        "dos convenciones a la vez, varios archivos generan el mismo identificador y el "
        "sistema se quedaría con uno solo de ellos en silencio. Por eso la función emite un "
        "aviso explícito cuando detecta la colisión, indicando qué archivos la producen.",
    )
    parrafo(doc, "Cada perfil devuelto contiene:")
    tabla(
        doc,
        ["Clave", "Significado"],
        [
            ("id", "Identificador del perfil, con el número del nombre del archivo."),
            ("num", "Número entero del perfil."),
            ("archivo", "Nombre exacto del archivo de slab, para trazabilidad."),
            ("lat_media", "Latitud media de la sección; se usa para ubicar los bordes norte y sur del conjunto."),
            ("slab", "Diccionario con la línea central reconstruida: distancias, longitud, latitud y profundidad."),
            ("topo_p, topo_alt", "Distancias y altitudes de la topografía associated al perfil."),
        ],
    )

    subseccion(doc, "7.3 Reconstrucción de la línea central")
    parrafo(
        doc,
        "La línea central de cada perfil no viene calculada: se reconstruye a partir de los "
        "puntos muestreados. El procedimiento ordena los puntos por su distancia p, los agrupa "
        "en bins de un kilómetro, y calcula en cada bin la media de longitud, de latitud y de "
        "profundidad absoluta, ignorando los valores NaN. El resultado es una polilínea de "
        "puntos con la que se proyectan los eventos.",
    )

    subseccion(doc, "7.4 Métricas calculadas")
    parrafo(
        doc,
        "La proyección es plana y local: no usa ni Mercator ni fórmulas esféricas. Se convierte "
        "la diferencia de longitud a kilómetros corrigiendo por el coseno de la latitud de "
        "referencia, que es la latitud media de la línea central del perfil. Para cada tramo de "
        "la polilínea se calcula el punto más cercano al evento, y de todos los tramos se toma "
        "el más cercano.",
    )
    tabla(
        doc,
        ["Métrica", "Definición", "Para qué sirve"],
        [
            ("perp_km", "Distancia perpendicular mínima del evento a la línea central.", "Medida de cercanía en planta; es la base de los topes laterales."),
            ("along_km", "Distancia a lo largo del perfil del punto más cercano, interpolada en el tramo.", "Permite consultar la profundidad de la placa en esa posición."),
            ("lat_proj", "Latitud del punto de la línea central más cercano al evento.", "Permite decidir si el evento quedó al norte o al sur del perfil."),
            ("exceso", "Cuánto cae el evento más allá del extremo del perfil, en km.", "Detecta eventos más allá del alcance longitudinal de la sección."),
            ("residuo_km", "Diferencia entre la profundidad del evento y la profundidad de la placa en along_km.", "Mide el desajuste vertical respecto del modelo."),
            ("dist_asoc", "Raíz de la suma del cuadrado de perp_km y del residuo ponderado por K.", "Métrica compuesta que decide a qué perfil se asigna el evento."),
        ],
    )
    parrafo(
        doc,
        "La profundidad de la placa se obtiene por interpolación lineal sobre la distancia "
        "along_km. Si el evento cae fuera del rango de la sección o la profundidad de la placa "
        "es NaN, el residuo no se puede calcular; en ese caso se trata como cero para no "
        "penalizar artificialmente la asociación.",
    )

    subseccion(doc, "7.5 Modelo de cobertura")
    parrafo(
        doc,
        "El modelo no se basa en un radio fijo alrededor de cada perfil, sino en la cobertura "
        "que el propio conjunto de archivos define. Hay tres situaciones distintas:",
    )
    vineta(
        doc,
        "Interior del conjunto: el evento se asigna al perfil que minimiza dist_asoc. No "
        "existe un tope de distancia en planta, de modo que no quedan zonas sin cubrir entre "
        "perfiles contiguos; en esas zonas siempre hay un perfil más cercano.",
    )
    vineta(
        doc,
        "Bordes del conjunto: el perfil más al norte y el más al sur del set, identificados "
        "por su latitud media, admiten hasta medio hueco mediano entre perfiles más el margen "
        "de borde y la tolerancia. Hacia el interior no se aplica este tope.",
    )
    vineta(
        doc,
        "Extremo longitudinal: si la proyección del evento cae más allá del inicio o del fin de "
        "la sección, más el margen y la tolerancia, se declara fuera de cobertura. Es el caso "
        "de los eventos oceánicos lejanos respecto de una sección costera.",
    )
    parrafo(
        doc,
        "Si está activo un tope lateral por distancia perpendicular, este actúa como un límite "
        "adicional del respaldo lateral. Finalmente, si la distancia de asociación supera el "
        "umbral global, se prefiere el perfil más cercano en planta en lugar de descartar el "
        "evento, siempre que la cobertura lo permita.",
    )
    parrafo(
        doc,
        "Cuando el evento queda fuera, el sistema conserva las métricas del perfil más cercano "
        "para que el análisis pueda aprovecharlas, pero no le asigna perfil, y registra el "
        "motivo. Los motivos se evalúan en este orden, de modo que el primero que corresponde "
        "es el que se informa:",
    )
    tabla(
        doc,
        ["Motivo", "Significado"],
        [
            ("fuera_cobertura_extremo", "El evento cae más allá del alcance longitudinal de la sección."),
            ("fuera_umbral", "Se superó el tope lateral por distancia perpendicular."),
            ("fuera_cobertura_norte", "El evento quedó al norte del perfil más septentrional, más allá del borde admitido."),
            ("fuera_cobertura_sur", "El evento quedó al sur del perfil más meridional, más allá del borde admitido."),
            ("fuera_cobertura", "Caso genérico de cobertura, sin correspondencia con los anteriores."),
            ("sin_perfiles", "No se detectó ningún perfil válido en la carpeta grillas."),
            ("sin_coordenadas", "El evento no tenía coordenadas utilizables."),
        ],
    )

    subseccion(doc, "7.6 Parámetros")
    tabla(
        doc,
        ["Constante", "Valor por defecto", "Efecto"],
        [
            ("UMBRAL_DIST_KM", "110.0", "Distancia de asociación máxima para aceptar un perfil; por encima, se prefiere el más cercano en planta."),
            ("UMBRAL_PERP_KM", "None", "Tope lateral por distancia perpendicular. En None no hay tope."),
            ("MARGEN_BORDE_KM", "0.0", "Margen en kilómetros que se añade al medio hueco mediano en los bordes y al alcance longitudinal."),
            ("_TOL_KM", "1.0", "Tolerancia de 1 km añadida a los topes, para absorber el redondeo de la proyección."),
            ("K_PESO_PROFUNDIDAD", "1.0", "Peso del residuo de profundidad frente a la distancia perpendicular en dist_asoc."),
            ("ANCHO_BIN_KM", "1.0", "Ancho del bin usado para reconstruir la línea central."),
            ("GRADO_KM_LAT / GRADO_KM_LON", "111.0", "Kilómetros por grado usados en la proyección local."),
        ],
    )
    parrafo(
        doc,
        "La distancia de asociación, el tope lateral, el margen de borde y el peso de la "
        "profundidad pueden modificarse desde la línea de comandos de generajson.py, con las "
        "opciones --umbral, --umbral-perp, --margen-borde y --k, sin tocar el código. La "
        "tolerancia, el ancho de bin y los kilómetros por grado son internos y se cambian en el "
        "propio módulo.",
    )

    subseccion(doc, "7.7 Ejemplos verificados")
    parrafo(
        doc,
        "Con el conjunto de 32 perfiles instalado y los parámetros por defecto, estos tres casos "
        "ilustran el comportamiento del modelo:",
    )
    tabla(
        doc,
        ["Entrada (lon, lat, prof)", "Resultado"],
        [
            ("-71.00, -25.00, 30 km", "Perfil P009; along 101.43 km, perp 2.37 km, residuo 8.62 km, dist_asoc 8.94 km. Asignado."),
            ("-76.00, -16.00, 30 km", "Sin perfil; perp 417.89 km. Motivo: fuera_cobertura_extremo."),
            ("-60.00, -30.00, 30 km", "Sin perfil; perp 536.19 km, dist_asoc 548.29 km. Motivo: fuera_cobertura_extremo."),
        ],
    )

    # ---------------------------------------------------------------- 8
    seccion(doc, "8. Clasificación de eventos sospechosos")
    parrafo(
        doc,
        "sismicidad.py decide si un evento puede estar mal localizado. La idea es cruzar dos "
        "familias de evidencia independientes: una geométrica, contra la placa, y otra "
        "estadística, contra la sismicidad histórica de la zona. Un evento solo se marca como "
        "sospechoso si se cumplen pruebas de ambas familias, lo que reduce los falsos positivos "
        "que aparecerían con una regla de disyunción simple.",
    )
    parrafo(
        doc,
        "El catálogo histórico se consulta únicamente como referencia: sus eventos nunca se "
        "clasifican. Además, se descartan automáticamente los registros históricos que están a "
        "menos de medio kilómetro del evento evaluado, para que la comparación no sea circular.",
    )
    parrafo(doc, "Pruebas disponibles y su estado por defecto:")
    tabla(
        doc,
        ["Grupo", "Prueba", "Activa", "Criterio"],
        [
            ("slab", "residuo_max_km", "Sí", "La diferencia entre la profundidad del evento y la de la placa supera 30 km."),
            ("slab", "perp_max_km", "No", "La distancia perpendicular al perfil supera 55 km."),
            ("histórica", "knn", "Sí", "La distancia al vecino histórico más cercano (k = 5) supera 15 km."),
            ("histórica", "densidad_local", "No", "Demasiados vecinos en una ventana de 20 km y 10 km de banda."),
            ("histórica", "ventana_aislamiento", "No", "Muy pocos vecinos en una ventana amplia de 45 km."),
            ("histórica", "desvio_mediana_prof", "No", "La profundidad se aparta de la mediana local en 20 km."),
        ],
    )
    parrafo(
        doc,
        "Las reglas se combinan de forma declarativa mediante una constante de configuración. "
        "La regla activa exige al menos una prueba del grupo de la placa y al menos una del "
        "grupo histórico. Si el archivo de catálogo no existe, la opción configurada deja el "
        "grupo histórico sin cumplir, con lo que ningún evento resulta sospechoso por esa vía.",
    )
    parrafo(
        doc,
        "Un evento sin perfil asignado se marca siempre como sospechoso, con independencia de "
        "las pruebas: si no se pudo determinar la profundidad de la placa bajo el evento, no hay "
        "base para descartarlo.",
    )

    # ---------------------------------------------------------------- 9
    seccion(doc, "9. Generación de JSON y CSV")
    parrafo(
        doc,
        "generajson.py lee el CSV de eventos, asigna el perfil de cada uno, calcula los "
        "sospechosos y escribe tres archivos. La lectura es posicional: el número de columna "
        "determina el significado del campo, y la última columna se interpreta como responsable "
        "en la fuente de Seisan o como la marca de percepción en eventquery. Un cambio en el "
        "orden de las columnas del CSV de entrada se manifestaría como datos mal interpretados, "
        "no como un error.",
    )
    tabla(
        doc,
        ["Archivo de salida", "Contenido"],
        [
            ("datos/eventos_<fuente>.json", "Lista de eventos con perfil, distancias y marca de sospechoso. Base del panel de mapas."),
            ("datos/sospechosos_<fuente>.csv", "Solo los eventos sospechosos, con las pruebas que se cumplieron."),
            ("datos/conteo_perfiles_<fuente>.json", "Totales, cobertura del conjunto de grillas y desglose de motivos sin perfil."),
        ],
    )
    parrafo(
        doc,
        "El archivo de conteo cumple dos funciones: informa al operador y alimenta los títulos "
        "del panel, que leen de ahí cuántos eventos hay en total y cuántos por perfil. Además "
        "registra la cobertura del conjunto y el desglose por motivo, que son la base para "
        "diagnosticar si el conjunto de grillas es adecuado para el catálogo.",
    )
    parrafo(doc, "En la consola, tras procesar, se informa:")
    vineta(doc, "La cobertura del conjunto: rango de latitud y longitud y número de perfiles detectados.")
    vineta(doc, "El desglose de eventos sin perfil, con una etiqueta legible por cada motivo.")
    vineta(
        doc,
        "Un aviso cuando hay eventos fuera de cobertura, recomendando un conjunto de "
        "mayor cobertura o ajustar el margen de borde.",
    )

    # ---------------------------------------------------------------- 10
    seccion(doc, "10. Ploteo e interacción")
    parrafo(
        doc,
        "plotear.py dibuja tres tipos de figura: el perfil de sección con su planta asociada, "
        "la planta ampliada de una zona, y el panel de análisis con una tabla de perfiles y un "
        "minimapa. Es importante saber que el módulo no guarda imágenes: todo se muestra en "
        "ventanas interactivas. Los únicos archivos que escribe son la bitácora de la sesión.",
    )
    parrafo(
        doc,
        "El perfil de sección muestra en el panel izquierdo la planta y en el derecho el corte "
        "vertical, con el relieve y la interfase de placa, y el fondo de sismicidad histórica "
        "proyectado desde el catálogo. El eje de profundidad tiene un piso que evita que la "
        "figura se aplaste, y el eje vertical se recorta por encima del nivel del mar.",
    )
    parrafo(doc, "Archivos que escribe el ploteo:")
    tabla(
        doc,
        ["Archivo", "Modo", "Función"],
        [
            ("ploteo/percibidos.txt", "append", "Registro de los eventos percibidos que el operador ha ido revisando."),
            ("ploteo/resumen_sesion.log", "sobrescritura", "Resumen de la sesión actual: eventos, sospechosos y percibidos vistos."),
            ("ploteo/resumen_historico.log", "append", "Acumulado histórico de todas las sesiones."),
        ],
    )
    parrafo(
        doc,
        "Al abrir un JSON, plotear agrupa los eventos en tres conjuntos: los que tienen un "
        "perfil presente en el conjunto de grillas actual, los que no tienen perfil, y los que "
        "apuntan a un perfil que ya no existe. Los terceros son la señal de que el JSON se "
        "generó con otro conjunto de grillas; el módulo los separa, avisa cuántos son y "
        "explica que conviene regenerar el JSON, en lugar de dibujarlos contra una sección "
        "distinta. Esta separación es lo que impide que un mapa presente eventos con una "
        "localización aparente pero falsa.",
    )
    parrafo(
        doc,
        "La interacción está pensada para revisar muchos eventos uno a uno. Al hacer clic en un "
        "evento en cualquiera de los dos paneles, se resalta en el otro, lo que permite "
        "contrastar planta y sección. Los gestos de arrastre se ignoran para no estorbar el "
        "zoom y el desplazamiento. El recuadro de referencia geográfica muestra la localidad "
        "más cercana con su distancia y rumbo.",
    )
    parrafo(
        doc,
        "Los eventos sin perfil se reparten en tres territorios, con límites definidos por "
        "constantes: Nacional, Insular (Isla de Pascua) y Antártico. Un evento que cae fuera de "
        "todos ellos se asigna al territorio más cercano y se avisa en consola, indicando "
        "aproximadamente a qué distancia quedó el punto.",
    )
    tabla(
        doc,
        ["Constante de ploteo", "Valor", "Efecto"],
        [
            ("FIG_SIZE / FIG_SIZE_NACIONAL", "(15, 7) / (7, 13)", "Tamaño de la figura de sección y de la planta nacional."),
            ("DPI", "100", "Resolución de las figuras."),
            ("PROF_MAX_KM", "250", "Piso del eje de profundidad."),
            ("MARGEN_PLANTA_GRADOS", "0.6", "Margen en grados alrededor de la zona de la planta."),
            ("HIST_MAX_PERP_PROF_KM", "40.0", "Distancia máxima a la que el fondo histórico se considera vinculado a la sección."),
            ("MAX_EVENTOS_ETIQUETA", "60", "Límite de etiquetas de id para no saturar la figura."),
            ("RELIEVE_LOCAL", "lon -97..-53, lat -62..-4", "Recorte del relieve regional, que tiene prioridad sobre el global."),
        ],
    )

    # ---------------------------------------------------------------- 11
    seccion(doc, "11. Módulos auxiliares")
    tabla(
        doc,
        ["Módulo", "Qué hace"],
        [
            ("proc_query_harz_2.py", "Normaliza el CSV de la consulta pública: ordena por tiempo, separa el estado de percepción, y produce una versión amplia de 16 columnas y otra reducida de 8. Si no hay datos, borra los archivos que había generado."),
            ("revisaselect.py", "Lee select.out y separa los eventos revisables de los excluidos. Un evento solo se considera revisable si tiene magnitud, dos coordenadas, exactamente dos referenciasinstrumentales y no está marcado como otro tipo de evento. Las líneas con problemas de formato se apartan."),
            ("revisacollect.py", "Construye el CSV maestro con una fila por evento. Separa los que no tienen estructura de fase y los que reutilizan la estación de referencia, lo deja asentado en reportes propios, y unifica los nombres de analista corrigiendo variantes por distancia de edición mínima."),
            ("compara.py", "Compara el catálogo público contra el select con dos filtros. Marca como no publicado el evento sin coincidencia temporal, y como no actualizado el que coincide pero difiere en posición o profundidad. Además aísla los de magnitud 2.5 o superior."),
            ("revisaexcluidos.py", "Normaliza el listado de excluidos y detecta las líneas anómalas, como una hora fuera de rango o segundos ausentes, que se apartan a un archivo de revisión humana en lugar de descartarse en silencio."),
            ("repetidos.py", "Detecta repeticiones dentro de cada fuente por separado, con los mismos dos filtros de ventana, y genera los cuatro reportes correspondientes."),
            ("repetidosexclu.py", "Cruza los excluidos contra el select completo con una ventana temporal fija de 6 segundos, e informa el tiempo de proceso."),
        ],
    )
    parrafo(
        doc,
        "Un detalle de mantenimiento en revisaexcluidos.py: al terminar intenta borrar un "
        "archivo temporal cuyo nombre difiere en una letra del que realmente crea "
        "revisaselect.py. El borrado falla de forma silenciosa, es decir, no provoca ningún "
        "error ni efecto observable, pero deja el archivo temporal en disco.",
    )
    parrafo(
        doc,
        "traequery.py y extraesc.sh son independientes del pipeline. El primero consulta el "
        "catálogo público por SSH de forma interactiva, validando las fechas y admitiendo "
        "filtros de rango. El segundo extrae eventos del servidor SeisComp mediante scquery, "
        "recreando el entorno del servidor en la sesión remota, y convierte la salida a CSV. "
        "Ambos son la vía de entrada de datos y no intervienen en el procesamiento.",
    )

    # ---------------------------------------------------------------- 12
    seccion(doc, "12. Tabla consolidada de umbrales y criterios")
    parrafo(
        doc,
        "Resumen de todos los valores que gobiernan los resultados del sistema, con su archivo "
        "de origen. Cualquier modificación debe hacerse en el archivo indicado.",
    )
    tabla(
        doc,
        ["Parámetro", "Valor", "Archivo"],
        [
            ("Ventana temporal, filtro amplio", "6 s", "compara.py, repetidos.py"),
            ("Ventana espacial, filtro amplio", "2.0°", "compara.py, repetidos.py"),
            ("Ventana temporal, filtro estricto", "3 s", "compara.py, repetidos.py"),
            ("Ventana espacial, filtro estricto", "1.0°", "compara.py, repetidos.py"),
            ("Magnitud mínima para no publicados", "2.5", "compara.py"),
            ("Ventana de repetidos en excluidos", "6 s", "repetidosexclu.py"),
            ("Distancia de asociación máxima", "110 km", "asigna_perfiles.py"),
            ("Tope lateral por perpendicular", "sin tope", "asigna_perfiles.py"),
            ("Margen de borde", "0 km", "asigna_perfiles.py"),
            ("Tolerancia de topes", "1 km", "asigna_perfiles.py"),
            ("Peso de la profundidad", "1.0", "asigna_perfiles.py"),
            ("Residuo de profundidad para sospechoso", "30 km", "sismicidad.py"),
            ("Distancia al vecino histórico (k = 5)", "15 km", "sismicidad.py"),
            ("Distancia mínima de exclusión del histórico", "0.5 km", "sismicidad.py"),
            ("Distancia de edición para unificar analistas", "1", "revisacollect.py"),
            ("Máximo de filas visibles en el visor de repetidos", "2000", "app.py"),
        ],
    )

    # ---------------------------------------------------------------- 13
    seccion(doc, "13. Datos de referencia")
    tabla(
        doc,
        ["Archivo", "Formato", "Uso"],
        [
            ("base_2020_2026.dat", "Columnas separadas por tabulador: fecha y hora, latitud, longitud, profundidad, magnitud, tipo de magnitud, marca de percepción y región.", "Catálogo histórico consultado por sismicidad.py y dibujado como fondo en plotear.py."),
            ("base_2020_2026.csv", "CSV con la misma información en formato de texto separado por comas.", "Equivalente del anterior; no lo usa ningún módulo."),
            ("base_2023_2026.dat / .csv", "Subconjunto del catálogo desde 2023.", "No lo usa ningún módulo; se mantiene como respaldo."),
            ("localidades.csv", "CSV con columnas Nombre, Lon, Lat.", "Rotulado de los mapas y cálculo de la localidad más cercana. Los nombres de las columnas deben ser exactos."),
        ],
    )

    # ---------------------------------------------------------------- 14
    seccion(doc, "14. Dependencias y entorno")
    parrafo(
        doc,
        "El entorno se fija por completo en requirements.txt, con versiones exactas. Las "
        "dependencias de mayor peso son Cartopy y sus proyecciones, Shapely para la geometría de "
        "las regiones, pyshp para leer los límites administrativos, adjustText para la "
        "colocación de las etiquetas de los eventos, Levenshtein para unificar nombres de "
        "analista, ttkbootstrap para el aspecto de la interfaz y SciPy para el cálculo de "
        "distancias.",
    )
    parrafo(
        doc,
        "Dos detalles del entorno que conviene conocer. El primero es el ya descrito: Shapely "
        "debe compilarse contra la GEOS del sistema para no entrar en conflicto con Cartopy. El "
        "segundo es que plotear desactiva el límite de tamaño de imagen de Pillow, porque los "
        "archivos de relieve tienen dimensiones que ese límite rechazaría por defecto.",
    )
    parrafo(
        doc,
        "Cartopy necesita además la descarga de sus bases de datos geográficas la primera vez "
        "que se usan, lo que requiere acceso a la red. Los límites de las regiones administrativas "
        "chilenas se obtienen de una fuente cartográfica en línea y, si no se pueden cargar, el "
        "sistema lo avisa y sigue dibujando sin ellos.",
    )

    # ---------------------------------------------------------------- 15
    seccion(doc, "15. Operación, diagnóstico y mantenimiento")
    subseccion(doc, "15.1 Cambiar el conjunto de grillas")
    parrafo(
        doc,
        "Para cambiar el conjunto de perfiles basta reemplazar los archivos de la carpeta "
        "grillas/ por el conjunto completo y volver a generar el JSON de la fuente "
        "correspondiente. No hay que tocar el código: los perfiles se detectan en cada "
        "generación y la cobertura se recalcula a partir de ellos.",
    )
    vineta(doc, "Verificar que la carpeta contenga solo una convención de nombres, y no una mezcla de formatos.")
    vineta(doc, "Comprobar la detección con la opción de listar perfiles, que informa identificador, extensión y si tiene topografía.")
    vineta(doc, "Regenerar el JSON; el panel se abrirá con el conjunto nuevo.")
    vineta(
        doc,
        "Si se mezclaron conjuntos, el conteo de perfiles y el número de eventos sin perfil del "
        "registro son la señal más rápida del problema.",
    )

    subseccion(doc, "15.2 Mensajes del sistema y qué significan")
    tabla(
        doc,
        ["Mensaje", "Causa", "Qué hacer"],
        [
            ("El JSON se generó con otro conjunto de grillas", "El JSON es anterior al conjunto de grillas actual.", "Regenerar el JSON de esa fuente. No corresponde revisar el mapa."),
            ("No se detectaron perfiles válidos en 'grillas'", "La carpeta está vacía o los archivos no siguen el patrón de nombres.", "Revisar el contenido y los nombres de grillas/; confirmar con la opción de listar."),
            ("Aviso de identificador producido por varios archivos", "La carpeta mezcla convenciones de nombres.", "Renombrar para dejar una sola convención."),
            ("Aviso: el conjunto de grillas no cubre todo el catálogo", "Hay eventos fuera del alcance del conjunto.", "Usar un conjunto de mayor cobertura, o revisar el margen de borde."),
            ("Aviso: evento fuera de los límites por defecto", "Un evento sin perfil quedó fuera de los tres territorios definidos.", "Informativo; el mapa se expande para incluirlo."),
            ("Aviso: no se pudieron cargar las regiones / localidades", "Faltan los archivos de límites o la red.", "Informativo; el resto del mapa se dibuja igual."),
        ],
    )

    subseccion(doc, "15.3 Límites conocidos")
    vineta(
        doc,
        "La profundidad de la placa solo se conoce dentro del rango longitudinal de cada "
        "sección; fuera de él no hay residuo que calcular.",
    )
    vineta(
        doc,
        "La proyección es plana y local, adecuada para el ámbito del país, pero no para "
        "distancias intercontinentales.",
    )
    vineta(
        doc,
        "El sistema no valida el orden de las columnas de los CSV de entrada: un cambio de "
        "formato se manifiesta como datos mal interpretados.",
    )
    vineta(
        doc,
        "No hay validación automática de que cada etapa haya producido los archivos que la "
        "siguiente espera; el usuario debe verificar los conteos del registro.",
    )

    ruta = os.path.join(RAIZ, "manual_tecnico.docx")
    doc.save(ruta)
    return ruta


# --------------------------------------------------------------------------
# MANUAL DE USUARIO
# --------------------------------------------------------------------------
def generar_usuario():
    doc = base_document()
    titulo(
        doc,
        "Manual de usuario — RCS",
        "Guía de operación diaria para la revisión de publicaciones sísmicas",
    )

    # ---------------------------------------------------------------- 1
    seccion(doc, "1. Qué hace el sistema")
    parrafo(
        doc,
        "RCS compara lo que publicó el sistema automático de detección sísmica con "
        "el catálogo público de eventos, y ayuda a revisar a mano los casos dudosos. En una "
        "sesión típica se hace lo siguiente:",
    )
    vineta(doc, "Se contrasta el catálogo público contra la selección del sistema para detectar eventos no publicados o publicados con datos distintos.")
    vineta(doc, "Se revisan los eventos que el sistema dejó fuera y los que aparecen repetidos.")
    vineta(doc, "Se genera un conjunto de eventos con su perfil de subducción asignado, y se revisan en un panel de mapas.")
    parrafo(
        doc,
        "El sistema no modifica nada de la fuente de datos: todos los resultados se escriben en "
        "archivos nuevos dentro de la carpeta de trabajo, y nada de lo que haga afecta a los "
        "sistemas de origen.",
    )

    # ---------------------------------------------------------------- 2
    seccion(doc, "2. Requisitos e instalación")
    parrafo(
        doc,
        "La aplicación se arranca con el script supervisor.sh, indicando el archivo de entrada "
        "y el archivo de salida temporal:",
    )
    codigo(doc, "./supervisor.sh <archivo_entrada.csv> <archivo_salida.dat>")
    vineta(doc, "La primera vez, el script crea el entorno virtual e instala las dependencias; puede tardar varios minutos.")
    vineta(doc, "Requiere Python 3.7 o superior, el paquete python3-venv, y las herramientas de compilación y la biblioteca GEOS del sistema.")
    vineta(
        doc,
        "Si la instalación falla al compilar Shapely, el mensaje indica instalar libgeos-dev y "
        "un compilador C. Es un paso conocido y necesario: Shapely debe compilarse en el "
        "sistema para que sea compatible con los mapas.",
    )
    vineta(
        doc,
        "Lance siempre el script desde la carpeta del proyecto. Los archivos de resultados se "
        "crean junto a donde se ejecuta, no junto al código.",
    )

    # ---------------------------------------------------------------- 3
    seccion(doc, "3. Recorrido A: análisis")
    parrafo(
        doc,
        "Con un archivo de entrada y un select.out ya disponibles, el botón de análisis ejecuta "
        "seis etapas en orden. El registro muestra el detalle de cada una y la barra indica el "
        "avance.",
    )
    tabla(
        doc,
        ["Etapa", "Qué hace", "Qué archivo deja para revisar"],
        [
            ("Procesando query", "Normaliza el archivo de entrada y ordena los eventos por hora.", "—"),
            ("Revisando select", "Decide qué eventos del select son revisables y cuáles quedan excluidos.", "informes/excluidos.txt"),
            ("Revisando collect", "Arma la tabla maestra de la revisión y un archivo por analista.", "datos/salida_collect.csv"),
            ("Comparando publicados v/s procesados", "Detecta eventos del catálogo público que no fueron publicados, y los que cambiaron.", "informes/no_pub_*.txt, informes/no_act_*.txt"),
            ("Revisando excluidos", "Normaliza el listado de excluidos y busca repeticiones.", "datos/excluidos.csv, informes/rep_seisan_exclu.txt"),
            ("Revisando repetidos", "Busca eventos repetidos dentro de cada fuente.", "informes/rep_*.txt"),
        ],
    )
    parrafo(
        doc,
        "Revise el registro al terminar: si alguna etapa se saltó, aparecerá anotado. Las "
        "etapas de comparación y de excluidos tienen condiciones, y si falta un archivo "
        "intermedio el sistema lo avisa en lugar de fallar.",
    )

    # ---------------------------------------------------------------- 4
    seccion(doc, "4. Recorrido B: generación de fuentes")
    parrafo(
        doc,
        "Terminado el análisis, los botones de fuente generan el conjunto de eventos con perfil "
        "y abren el panel de mapas. Hay tres botones disponibles:",
    )
    tabla(
        doc,
        ["Botón", "Qué genera", "Cuándo usarlo"],
        [
            ("Datos de Seisan", "El conjunto de eventos de la revisión actual, con perfil y marca de sospechoso.", "Es la fuente principal de trabajo."),
            ("Datos de eventquery", "Lo mismo, pero tomado del catálogo público completo.", "Para contrastar contra el catálogo público."),
            ("No publicados de Seisan", "El conjunto de eventos no publicados detectados en el paso de comparación.", "Para revisar específicamente lo que no se publicó."),
        ],
    )
    parrafo(
        doc,
        "Un detalle importante: el botón de no publicados escribe sobre el mismo archivo que el "
        "botón de Seisan. Si necesita conservar ambos conjuntos, genere el de Seisan, revise, y "
        "después genere el de no publicados.",
    )
    parrafo(
        doc,
        "Al terminar cada generación, el registro informa cuántos eventos quedaron con perfil "
        "y cuántos sin perfil, con el detalle de los motivos. Es el indicador principal para "
        "saber si el conjunto de grillas es adecuado para los datos con los que está trabajando.",
    )

    # ---------------------------------------------------------------- 5
    seccion(doc, "5. Archivos generados")
    parrafo(
        doc,
        "Todo se escribe dentro de la carpeta desde la que se lanzó la aplicación. Las "
        "carpetas se crean solas la primera vez.",
    )
    tabla(
        doc,
        ["Carpeta", "Archivo", "Qué contiene"],
        [
            ("datos/", "salida_collect.csv", "Tabla maestra de la revisión, una fila por evento."),
            ("datos/", "new_2_<archivo>.csv", "Eventos de entrada ya normalizados, 8 columnas."),
            ("datos/", "excluidos.csv", "Eventos excluidos, con su fecha y responsable."),
            ("datos/", "eventos_<fuente>.json", "Eventos con perfil asignado; es lo que alimenta los mapas."),
            ("datos/", "sospechosos_<fuente>.csv", "Solo los eventos marcados como sospechosos, con las pruebas que se cumplieron."),
            ("datos/", "conteo_perfiles_<fuente>.json", "Totales, cobertura de las grillas y motivos de los eventos sin perfil."),
            ("informes/", "no_pub_*.txt / no_act_*.txt", "Eventos no publicados y no actualizados, con los filtros amplio y estricto."),
            ("informes/", "rep_*.txt", "Repeticiones detectadas en cada fuente y con cada filtro."),
            ("informes/", "excluidos.txt / revisar.txt", "Exclusiones normalizadas y las líneas que requieren revisión manual."),
            ("analistas/", "<analista>.csv", "Eventos correspondientes a cada analista."),
            ("ploteo/", "percibidos.txt", "Registro de los eventos percibidos que ha revisado."),
            ("ploteo/", "resumen_sesion.log", "Resumen de la sesión: eventos, sospechosos y percibidos vistos."),
        ],
    )

    # ---------------------------------------------------------------- 6
    seccion(doc, "6. Cómo leer los resultados")
    subseccion(doc, "6.1 Eventos no publicados y no actualizados")
    parrafo(
        doc,
        "La comparación se hace con dos filtros, y ambos se generan siempre. El filtro amplio "
        "considera que dos eventos son el mismo si están a menos de 6 segundos y dentro de 2 "
        "grados; el estricto, si están a menos de 3 segundos y dentro de 1 grado. El amplio "
        "encuentra más coincidencias y por tanto deja menos eventos marcados como no publicados; "
        "el estricto es más exigente. Lo habitual es mirar el amplio primero y usar el estricto "
        "para los casos que importan.",
    )
    parrafo(
        doc,
        "Un evento no actualizado es aquel que sí se publicó pero difiere de lo procesado, en "
        "posición o profundidad. Y entre los no publicados se aíslan los de magnitud 2.5 o "
        "superior, que son los que suelen justificar una revisión con más detalle.",
    )

    subseccion(doc, "6.2 Repetidos")
    parrafo(
        doc,
        "La detección de repeticiones se hace por separado dentro de cada fuente, con los mismos "
        "dos filtros, y genera cuatro archivos: dos para la fuente pública y dos para la de "
        "Seisan. La ventana de búsqueda es de 6 segundos en el filtro amplio y de 3 en el "
        "estricto, dentro del margen espacial correspondiente.",
    )
    parrafo(
        doc,
        "La ventana del visor de repetidos está limitada a 2000 filas para que la aplicación no "
        "se congele con archivos grandes. Si hay más eventos, la propia ventana indica cuántos "
        "quedan fuera de la vista, y el archivo completo sigue en informes/.",
    )

    subseccion(doc, "6.3 Excluidos y revisión manual")
    parrafo(
        doc,
        "Un evento queda excluido cuando al sistema le faltan piezas: la magnitud, las dos "
        "coordenadas, o las referencias instrumentales esperadas. La lista de excluidos se "
        "normaliza y, si alguna línea tiene datos raros, como una hora imposible o segundos "
        "ausentes, no se descarta en silencio: se escribe en un archivo aparte para que la "
        "revisen a mano.",
    )

    subseccion(doc, "6.4 Sospechosos de mala localización")
    parrafo(
        doc,
        "El sistema marca como sospechoso un evento cuando hay indicios en dos frentes a la "
        "vez: que su profundidad no calza con la placa que debería tener debajo, y que no hay "
        "sismicidad histórica comparable cerca. Exigir las dos cosas es deliberado, porque por "
        "separado cada criterio produce demasiados falsos positivos. Los criterios están "
        "configurados con una holgura de 30 km en profundidad y de 15 km al vecino histórico más "
        "cercano.",
    )
    parrafo(
        doc,
        "En el archivo de sospechosos, la columna de criterios indica qué pruebas se cumplieron "
        "en cada caso, lo que permite entender por qué está en la lista.",
    )

    subseccion(doc, "6.5 Eventos sin perfil: qué significa")
    parrafo(
        doc,
        "Un evento puede quedarse sin perfil asignado, y no siempre es un error. El sistema "
        "distingue el motivo y lo informa en el registro y en el conteo de perfiles:",
    )
    tabla(
        doc,
        ["Motivo", "Qué significa en la práctica"],
        [
            ("Fuera de cobertura (norte o sur del conjunto)", "El evento quedó más allá del último perfil de esa dirección, en una zona que ninguna sección del conjunto cubre."),
            ("Fuera del alcance de las secciones", "El evento está más allá del comienzo o del final de las secciones, típicamente en el océano lejano."),
            ("Fuera del tope lateral", "Solo ocurre si se fijó un tope lateral explícito; con la configuración normal no aparece."),
            ("Sin coordenadas", "El evento venía sin datos suficientes para ubicarlo."),
        ],
    )
    parrafo(
        doc,
        "Estos eventos se dibujan aparte, en el mapa de sin perfil, y no se mezclan con los de "
        "las secciones. Todos ellos quedan marcados como sospechosos por definición, porque sin "
        "perfil no hay forma de contrastar la profundidad. Si la proporción de eventos sin perfil "
        "es alta, lo habitual es que el conjunto de grillas no cubra la zona de trabajo, y la "
        "solución es usar un conjunto con mayor cobertura antes que revisar caso por caso.",
    )

    # ---------------------------------------------------------------- 7
    seccion(doc, "7. El panel de análisis")
    parrafo(
        doc,
        "Al generar una fuente se abre el panel de análisis. A la izquierda hay una tabla con "
        "los perfiles encontrados, y a la derecha un perfil de referencia.",
    )
    tabla(
        doc,
        ["Columna", "Qué indica"],
        [
            ("n", "Cantidad de eventos del perfil."),
            ("perfil", "Identificador del perfil."),
            ("eventos", "Total de eventos asignados a ese perfil."),
            ("sosp", "Cuántos de ellos están marcados como sospechosos."),
            ("perc", "Cuántos fueron percibidos."),
            ("along", "Posición a lo largo del perfil, en kilómetros."),
        ],
    )
    parrafo(doc, "Para revisar un perfil:")
    vineta(doc, "Con un solo clic en la fila, el perfil de referencia de la derecha se redibuja con ese perfil; sirve para tantear sin abrir ventanas.")
    vineta(doc, "Con doble clic en la fila, o con el botón Abrir detalle, se abre la ventana completa del perfil.")
    vineta(doc, "Al hacer clic en un evento en el corte, se resalta también en la planta, y viceversa. Así se revisan posición y profundidad a la vez.")
    vineta(doc, "El recuadro de referencia indica la localidad más cercana, con su distancia y rumbo.")
    vineta(doc, "Use el botón Detener para interrumpir el despliegue de mapas cuando ya tenga suficiente información.")
    parrafo(
        doc,
        "Los eventos sin perfil se reparten en tres mapas según el territorio: Nacional, "
        "Insular para Isla de Pascua, y Antártico. Si algún evento queda fuera de los tres, el "
        "sistema lo asigna al más cercano y lo avisa; el mapa se expande para incluirlo.",
    )
    parrafo(
        doc,
        "La navegación es la habitual de los mapas: zoom y desplazamiento con los controles, y "
        "el conmutador de modo con la tecla Escape. Los eventos percebidos se distinguen por su "
        "color, y los sospechosos llevan un contorno marcado.",
    )

    # ---------------------------------------------------------------- 8
    seccion(doc, "8. Parámetros que se pueden ajustar")
    parrafo(
        doc,
        "El sistema tiene algunos parámetros pensados para ajustarse, aunque la configuración "
        "normal no requiere tocarlos. Conviene conocerlos antes de modificar cualquiera.",
    )
    tabla(
        doc,
        ["Parámetro", "Qué hace", "Cuándo tocarlo"],
        [
            ("Margen de borde", "Amplía hasta dónde se considera que el conjunto de grillas cubre su borde y el alcance de las secciones.", "Si el conjunto cubre algo más de lo que el sistema asume y hay eventos válidos marcados como fuera de cobertura."),
            ("Tope lateral", "Impone una distancia máxima al perfil más cercano.", "Solo en estudios muy específicos; con la configuración normal está desactivado."),
            ("Distancia de asociación máxima", "Límite de distancia para aceptar la asociación con un perfil.", "Casos poco frecuentes; preferir otro conjunto de grillas."),
            ("Peso de la profundidad", "Cuánto pesa el desajuste de profundidad frente a la distancia en planta al decidir el perfil.", "Solo si conoce la respuesta correcta para varios eventos de referencia."),
            ("Criterios de sospechoso", "Los umbrales de las pruebas que marcan eventos como sospechosos.", "Con análisis de la tasa de confirmación; un umbral más alto marca menos eventos."),
            ("Ventanas de comparación y de repeticiones", "Los filtros de 6 y 3 segundos con sus márgenes espaciales.", "Rara vez; el filtro amplio y el estricto ya cubren los dos extremos."),
        ],
    )
    parrafo(
        doc,
        "El margen de borde, el tope lateral, la distancia de asociación máxima y el peso de "
        "la profundidad se pueden pasar como argumentos al generar el JSON, sin modificar el "
        "código. Los criterios de sospechoso se definen en un diccionario de configuración "
        "dentro de sismicidad.py, y las ventanas de comparación, en las constantes de cada "
        "módulo. Si cambia alguno de estos últimos, regenere el JSON para que el cambio tenga "
        "efecto.",
    )

    # ---------------------------------------------------------------- 9
    seccion(doc, "9. Problemas frecuentes")
    tabla(
        doc,
        ["Síntoma", "Causa probable", "Qué hacer"],
        [
            ("El JSON se generó con otro conjunto de grillas", "Se cambió la carpeta de grillas después de generar el JSON.", "Vuelva a generar el JSON de esa fuente. No revise el mapa: los perfiles ya no corresponden."),
            ("No se detectaron perfiles válidos", "La carpeta de grillas está vacía o los archivos no siguen el patrón de nombres.", "Revise los nombres de los archivos; deben llamarse slabP###.tmp con su topografía correspondiente."),
            ("Aviso de que un identificador lo producen varios archivos", "La carpeta mezcla dos convenciones de nombres.", "Deje una sola convención en la carpeta y vuelva a generar."),
            ("El conjunto de grillas no cubre todo el catálogo", "Muchos eventos quedan fuera del alcance de las secciones.", "Use un conjunto de mayor cobertura, o evalúe ampliar el margen de borde."),
            ("El panel no abre con la lista de perfiles", "El JSON no existe todavía o la carpeta de grillas está vacía.", "Genere primero la fuente; compruebe que la carpeta de grillas tenga contenido."),
            ("Falla la instalación de dependencias", "Falta la biblioteca GEOS del sistema o un compilador.", "Instale libgeos-dev y un compilador C, y vuelva a ejecutar el arranque."),
            ("Faltan los límites de las regiones en el mapa", "No se pudieron descargar los datos cartográficos.", "Revise la conexión a la red; el resto del mapa se dibuja igual."),
            ("La lista de repetidos se ve truncada", "El visor muestra como máximo 2000 filas por diseño.", "Abra el archivo completo en informes/; la ventana indica cuántos faltan."),
        ],
    )

    # ---------------------------------------------------------------- 10
    seccion(doc, "10. Cambiar el conjunto de grillas")
    parrafo(
        doc,
        "Es la operación más frecuente y no requiere tocar el código. El sistema reconoce los "
        "archivos que encuentra en la carpeta cada vez que genera el JSON.",
    )
    vineta(doc, "Copie el conjunto completo de perfiles a la carpeta grillas/, reemplazando el anterior.")
    vineta(doc, "Revise que la carpeta contenga una sola convención de nombres, sin mezclar archivos de formatos distintos.")
    vineta(doc, "Verifique la detección antes de continuar: el sistema lista los perfiles encontrados con su extensión y si tienen topografía.")
    vineta(doc, "Vuelva a generar el JSON de la fuente que iba a revisar.")
    vineta(
        doc,
        "Compare en el registro el número de eventos con perfil antes y después. Si la mejora "
        "es la esperada, el conjunto nuevo cubre mejor la zona.",
    )
    parrafo(
        doc,
        "Si dejó abierto un panel con un JSON del conjunto anterior, no lo revise: los perfiles "
        "que le dibujen no corresponden. El propio sistema lo advierte y separa esos eventos.",
    )

    # ---------------------------------------------------------------- 11
    seccion(doc, "11. Bitácora de la sesión")
    parrafo(
        doc,
        "El sistema deja registro de cada sesión en la carpeta ploteo/, útil para retomar el "
        "trabajo o para saber qué se revisó.",
    )
    tabla(
        doc,
        ["Archivo", "Qué registra"],
        [
            ("resumen_sesion.log", "Una línea por sesión: fecha y hora, fuente, total de eventos, cuántos sospechosos y cuántos percibidos vio."),
            ("resumen_historico.log", "El mismo registro acumulado de todas las sesiones."),
            ("percibidos.txt", "Los eventos percibidos que ha ido revisando, para no volver a verlos."),
        ],
    )
    parrafo(
        doc,
        "El resumen de la sesión se sobrescribe en cada ejecución, mientras que el histórico y "
        "el registro de percibidos se van acumulando. Si quiere conservar el detalle de una "
        "sesión concreta, copie el archivo de resumen antes de la siguiente.",
    )

    # ---------------------------------------------------------------- 12
    seccion(doc, "12. Resumen de mensajes del sistema")
    tabla(
        doc,
        ["Mensaje en pantalla o registro", "Qué significa"],
        [
            ("cobertura del set: lat ... | lon ...", "Rango geográfico que cubren las grillas instaladas y cuántos perfiles se detectaron."),
            ("Eventos sin perfil, con su motivo", "Cuántos eventos quedaron fuera y por qué. Es la lectura más útil después de generar un JSON."),
            ("Aviso: el set de grillas no cubre todo el catálogo", "Hay eventos fuera de la cobertura del conjunto."),
            ("Aviso: el JSON se generó con otro conjunto de grillas", "El JSON es anterior a las grillas actuales; sus eventos no se dibujan contra las secciones nuevas."),
            ("Aviso: el id 'Pxxx' lo producen varios archivos", "Se mezclaron dos convenciones de nombres en la carpeta de grillas."),
            ("Aviso: evento fuera de los límites por defecto", "Un evento sin perfil quedó fuera de los tres territorios; se asignó al más cercano."),
            ("Aviso: no se pudieron cargar las regiones / localidades", "Faltan los datos cartográficos o no hay red. El mapa se dibuja igual, sin esos extras."),
        ],
    )
    parrafo(
        doc,
        "Los dos primeros mensajes no son errores: son el resumen del resultado. El resto son "
        "avisos que conviene leer antes de dar por buena una revisión.",
    )

    ruta = os.path.join(RAIZ, "manual_usuario.docx")
    doc.save(ruta)
    return ruta


if __name__ == "__main__":
    print(generar_tecnico())
    print(generar_usuario())
